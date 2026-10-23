"""A paragraph that carries a redundant nested w declaration, removed locally after a first save."""
import io, sys, tempfile, pathlib, docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import rdocx

folder = pathlib.Path(tempfile.mkdtemp())
d = docx.Document()
d.add_paragraph("first")
d.element.body.insert(1, parse_xml(f'<w:p {nsdecls("w")}><w:r><w:t>owner</w:t></w:r></w:p>'))
d.add_paragraph("last")
d.save(folder / "a.docx")
for save_first in (False, True):
    doc = rdocx.Document.open(str(folder / "a.docx"))
    if save_first:
        doc.save(str(folder / "b.docx"))
    try:
        doc.remove_content(1)
        doc.save(str(folder / "c.docx"))
        print("save_first", save_first, "ok", [p.text for p in rdocx.Document.open(str(folder / "c.docx")).paragraphs])
    except Exception as error:
        print("save_first", save_first, "ERROR", error)
